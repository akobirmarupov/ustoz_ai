"""Tayyor .pptx va .docx fayllarni yasash."""
import random
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches
from pptx.util import Pt as PPt

LABELS = {
    "uz": {"intro": "KIRISH", "conclusion": "XULOSA", "refs": "FOYDALANILGAN ADABIYOTLAR", "toc": "MUNDARIJA",
           "topic": "Mavzu", "by": "Bajardi", "checked": "Tekshirdi", "annotation": "Annotatsiya",
           "keywords": "Kalit so'zlar", "thanks": "E'tiboringiz uchun rahmat!",
           "referat": "REFERAT", "mustaqil": "MUSTAQIL ISH", "kurs": "KURS ISHI", "esse": "ESSE"},
    "ru": {"intro": "ВВЕДЕНИЕ", "conclusion": "ЗАКЛЮЧЕНИЕ", "refs": "СПИСОК ЛИТЕРАТУРЫ", "toc": "СОДЕРЖАНИЕ",
           "topic": "Тема", "by": "Выполнил(а)", "checked": "Проверил(а)", "annotation": "Аннотация",
           "keywords": "Ключевые слова", "thanks": "Спасибо за внимание!",
           "referat": "РЕФЕРАТ", "mustaqil": "САМОСТОЯТЕЛЬНАЯ РАБОТА", "kurs": "КУРСОВАЯ РАБОТА", "esse": "ЭССЕ"},
    "en": {"intro": "INTRODUCTION", "conclusion": "CONCLUSION", "refs": "REFERENCES", "toc": "CONTENTS",
           "topic": "Topic", "by": "Prepared by", "checked": "Checked by", "annotation": "Abstract",
           "keywords": "Keywords", "thanks": "Thank you for your attention!",
           "referat": "ESSAY", "mustaqil": "INDEPENDENT WORK", "kurs": "COURSE PAPER", "esse": "ESSAY"},
}


# ---------------- PowerPoint ----------------
#
# Har bir taqdimot uchun tasodifiy uslub tanlanadi: rang palitrasi + muqova maketi +
# slayd maketi + shriftlar + marker. Shu sababli bir xil mavzudagi slaydlar ham
# har xil ko'rinishda chiqadi.

def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


# primary — asosiy to'q rang, accent/accent2 — urg'u ranglari, bg — slayd foni,
# card — kartochka foni, ink — matn, heading — sarlavha, light — to'q fondagi matn.
PALETTES = [
    dict(primary="0F1F4B", accent="00C2FF", accent2="F2B632", bg="FFFFFF", card="F1F5FB",
         ink="262F40", heading="0F1F4B", muted="8A93A6", light="FFFFFF", soft="C9D6F0"),
    dict(primary="1B4332", accent="52B788", accent2="F4A261", bg="F8F9F4", card="E9F2EC",
         ink="23302A", heading="1B4332", muted="7F8C84", light="FFFFFF", soft="CFE6D8"),
    dict(primary="5C1A2B", accent="E07A5F", accent2="F2CC8F", bg="FFF8F0", card="F8EBE0",
         ink="3A2A2E", heading="5C1A2B", muted="9C8A8E", light="FFFFFF", soft="EED3D9"),
    dict(primary="1E1B2E", accent="A78BFA", accent2="F472B6", bg="1E1B2E", card="2C2842",
         ink="E5E7EB", heading="FFFFFF", muted="9B97B5", light="FFFFFF", soft="CFC8F5"),
    dict(primary="0F4C5C", accent="FB8B24", accent2="E36414", bg="FFFFFF", card="EEF5F6",
         ink="22333A", heading="0F4C5C", muted="83979C", light="FFFFFF", soft="BFDCE3"),
    dict(primary="222831", accent="FFD369", accent2="00ADB5", bg="EEEEEE", card="FFFFFF",
         ink="2B313A", heading="222831", muted="8C9199", light="FFFFFF", soft="D7DBE0"),
    dict(primary="1D3557", accent="E63946", accent2="A8DADC", bg="F1FAEE", card="FFFFFF",
         ink="1D3557", heading="1D3557", muted="7C8BA1", light="FFFFFF", soft="C8DDEE"),
    dict(primary="2B2D42", accent="EF8354", accent2="BFC0C0", bg="2B2D42", card="383B54",
         ink="EDEEF2", heading="FFFFFF", muted="A3A6B8", light="FFFFFF", soft="D9DAE3"),
    dict(primary="3D2C8D", accent="FF9F1C", accent2="2EC4B6", bg="FFFFFF", card="F3F1FB",
         ink="2D2640", heading="3D2C8D", muted="8E88A3", light="FFFFFF", soft="D6D0F2"),
    dict(primary="003049", accent="F77F00", accent2="FCBF49", bg="FDF6EC", card="FFFFFF",
         ink="1F2E38", heading="003049", muted="83919A", light="FFFFFF", soft="BCD3E0"),
    dict(primary="0B3D2E", accent="D4A373", accent2="FAEDCD", bg="0B3D2E", card="14523F",
         ink="EEF3EF", heading="FAEDCD", muted="9DB8AD", light="FFFFFF", soft="D6E6DE"),
    dict(primary="B23A48", accent="2F3E46", accent2="F6BD60", bg="FFFFFF", card="FBEFF0",
         ink="2F3E46", heading="B23A48", muted="8D979C", light="FFFFFF", soft="F4D3D7"),
]

# (sarlavha shrifti, matn shrifti) — Office/WPS'da odatda mavjud shriftlar
FONTS = [
    ("Calibri", "Calibri"), ("Georgia", "Calibri"), ("Arial", "Arial"),
    ("Trebuchet MS", "Trebuchet MS"), ("Cambria", "Calibri"), ("Verdana", "Verdana"),
    ("Century Gothic", "Century Gothic"), ("Segoe UI", "Segoe UI"),
]
MARKERS = ["■", "●", "▸", "◆", "—", "✓", "num"]
COVERS = ["full", "split", "band", "corner"]
LAYOUTS = ["topbar", "sidebar", "header", "split", "card"]

W, H = 13.333, 7.5


class _Theme:
    def __init__(self, rnd):
        pal = rnd.choice(PALETTES)
        for k, v in pal.items():
            setattr(self, k, _rgb(v))
        if pal["bg"] == pal["primary"]:
            # qorong'i fon: panellar fonga qo'shilib ketmasligi uchun ochroq rang
            self.primary = self.card
        self.head_font, self.body_font = rnd.choice(FONTS)
        self.marker = rnd.choice(MARKERS)
        self.cover = rnd.choice(COVERS)
        self.layout = rnd.choice(LAYOUTS)


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _shape(slide, kind, x, y, w, h, color):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _rect(slide, x, y, w, h, color):
    return _shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, color)


def _circle(slide, x, y, d, color):
    return _shape(slide, MSO_SHAPE.OVAL, x, y, d, d, color)


def _text(slide, x, y, w, h, text, size, color, bold=False, anchor=MSO_ANCHOR.TOP,
          align=PP_ALIGN.LEFT, font=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if font:
            run.font.name = font
    return box


def _bullets(slide, t: _Theme, x, y, w, h, bullets, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for j, b in enumerate(bullets):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_after = PPt(10)
        mark = p.add_run()
        if t.marker == "num":
            mark.text = f"{j + 1:02d}   "
            mark.font.bold = True
            mark.font.size = PPt(size - 2)
        else:
            mark.text = f"{t.marker}  "
            mark.font.size = PPt(size - 6)
        mark.font.color.rgb = t.accent
        mark.font.name = t.head_font
        run = p.add_run()
        run.text = b
        run.font.size = PPt(size)
        run.font.color.rgb = t.ink
        run.font.name = t.body_font


def _title_size(title: str, base: int) -> int:
    return base if len(title) <= 45 else base - 4 if len(title) <= 80 else base - 8


# ---- muqova (birinchi slayd) ----

def _cover(prs, t: _Theme, title: str, titul: list[str]):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    info = "\n".join(titul)
    size = _title_size(title, 40)
    if t.cover == "full":
        _bg(s, t.primary)
        _circle(s, 10.2, -1.6, 4.6, t.accent)
        _circle(s, 11.9, 5.4, 2.6, t.accent2)
        _rect(s, 0.8, 4.05, 1.6, 0.07, t.accent)
        _rect(s, 2.5, 3.98, 0.2, 0.2, t.accent2)
        _text(s, 0.8, 1.2, 9.2, 2.7, title, size, t.light, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        if info:
            _text(s, 0.8, 4.4, 9.2, 2.4, info, 18, t.soft, font=t.body_font)
    elif t.cover == "split":
        _bg(s, t.bg)
        _rect(s, 0, 0, 4.6, H, t.primary)
        _circle(s, 1.3, 2.75, 2.0, t.accent)
        _circle(s, 2.6, 4.2, 0.9, t.accent2)
        _rect(s, 5.3, 4.05, 1.4, 0.07, t.accent)
        _text(s, 5.3, 0.9, 7.4, 3.0, title, size - 2, t.heading, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        if info:
            _text(s, 5.3, 4.4, 7.4, 2.6, info, 17, t.ink, font=t.body_font)
    elif t.cover == "band":
        _bg(s, t.bg)
        _rect(s, 0, 1.3, W, 3.1, t.primary)
        _rect(s, 0, 4.4, W, 0.1, t.accent)
        _circle(s, 11.4, 0.5, 1.2, t.accent2)
        _text(s, 0.8, 1.4, 11.7, 2.9, title, size, t.light, bold=True, anchor=MSO_ANCHOR.MIDDLE,
              align=PP_ALIGN.CENTER, font=t.head_font)
        if info:
            _text(s, 0.8, 4.8, 11.7, 2.4, info, 17, t.ink, align=PP_ALIGN.CENTER, font=t.body_font)
    else:  # corner
        _bg(s, t.bg)
        tri = _shape(s, MSO_SHAPE.RIGHT_TRIANGLE, 8.3, 2.5, 5.033, 5.0, t.primary)
        tri._element.spPr.xfrm.set("flipH", "1")  # to'g'ri burchak o'ng pastda
        _circle(s, -1.2, -1.2, 3.2, t.accent)
        _circle(s, 1.6, 0.4, 0.6, t.accent2)
        _rect(s, 0.8, 4.25, 2.2, 0.08, t.accent)
        _text(s, 0.8, 1.4, 9.0, 2.7, title, size, t.heading, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        if info:
            _text(s, 0.8, 4.6, 8.0, 2.4, info, 17, t.ink, font=t.body_font)


# ---- asosiy slaydlar ----

def _content(prs, t: _Theme, i: int, total: int, sl: dict):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, t.bg)
    bullets = sl["bullets"][:7]
    size = 20 if len(bullets) <= 5 else 18
    title = sl["title"]
    tsize = _title_size(title, 30)
    num = f"{i}/{total}"

    if t.layout == "topbar":
        _rect(s, 0, 0, W, 0.14, t.primary)
        _text(s, 0.8, 0.45, 11.7, 1.0, title, tsize, t.heading, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        _rect(s, 0.8, 1.55, 1.1, 0.06, t.accent)
        _bullets(s, t, 0.8, 1.9, 11.7, 4.9, bullets, size)
        _text(s, 11.6, 6.85, 1.2, 0.4, num, 12, t.muted, align=PP_ALIGN.RIGHT, font=t.body_font)
        _rect(s, 12.95, 6.95, 0.12, 0.12, t.accent2)
    elif t.layout == "sidebar":
        _rect(s, 0, 0, 0.35, H, t.primary)
        _rect(s, 0.35, 0, 0.08, H, t.accent)
        _text(s, 11.0, 0.35, 1.9, 0.9, f"{i:02d}", 40, t.accent2, bold=True, align=PP_ALIGN.RIGHT, font=t.head_font)
        _text(s, 1.0, 0.45, 10.0, 1.1, title, tsize, t.heading, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        _rect(s, 1.0, 1.65, 11.8, 0.03, t.muted)
        _bullets(s, t, 1.0, 1.95, 11.6, 4.9, bullets, size)
        _text(s, 11.6, 6.9, 1.2, 0.4, num, 12, t.muted, align=PP_ALIGN.RIGHT, font=t.body_font)
    elif t.layout == "header":
        _rect(s, 0, 0, W, 1.55, t.primary)
        _rect(s, 0, 1.55, W, 0.08, t.accent)
        _circle(s, 12.2, 0.45, 0.65, t.accent2)
        _text(s, 0.8, 0.2, 11.1, 1.2, title, tsize - 2, t.light, bold=True, anchor=MSO_ANCHOR.MIDDLE, font=t.head_font)
        _bullets(s, t, 0.8, 2.0, 11.7, 4.8, bullets, size)
        _text(s, 0.8, 6.9, 3.0, 0.4, num, 12, t.muted, font=t.body_font)
    elif t.layout == "split":
        _rect(s, 0, 0, 4.4, H, t.primary)
        _circle(s, 3.3, 5.9, 2.2, t.accent)
        _text(s, 0.6, 0.6, 2.5, 0.8, f"{i:02d}", 36, t.accent2, bold=True, font=t.head_font)
        _rect(s, 0.6, 1.5, 0.9, 0.06, t.accent)
        _text(s, 0.6, 1.8, 3.4, 4.0, title, min(tsize, 26), t.light, bold=True, font=t.head_font)
        _bullets(s, t, 5.0, 0.8, 7.7, 6.0, bullets, size - 1)
        _text(s, 11.6, 6.9, 1.2, 0.4, num, 12, t.muted, align=PP_ALIGN.RIGHT, font=t.body_font)
    else:  # card
        _text(s, 0.8, 0.35, 11.7, 1.0, title, tsize, t.heading, bold=True, anchor=MSO_ANCHOR.BOTTOM, font=t.head_font)
        _rect(s, 0.8, 1.45, 0.6, 0.06, t.accent)
        _rect(s, 1.45, 1.45, 0.25, 0.06, t.accent2)
        card = _shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 1.8, 11.73, 5.0, t.card)
        card.adjustments[0] = 0.05
        _rect(s, 0.8, 2.1, 0.09, 4.4, t.accent)
        _bullets(s, t, 1.2, 2.05, 11.1, 4.6, bullets, size)
        _text(s, 11.6, 6.95, 1.2, 0.4, num, 12, t.muted, align=PP_ALIGN.RIGHT, font=t.body_font)


# ---- yakuniy slayd ----

def _closing(prs, t: _Theme, text: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if t.cover in ("full", "band"):
        _bg(s, t.primary)
        _circle(s, -1.5, 4.6, 4.0, t.accent)
        _circle(s, 11.6, 0.6, 1.4, t.accent2)
        color = t.light
    else:
        _bg(s, t.bg)
        _rect(s, 0, H - 0.5, W, 0.5, t.primary)
        _circle(s, 10.8, -1.4, 3.6, t.accent)
        _circle(s, 0.9, 5.2, 0.8, t.accent2)
        color = t.heading
    _text(s, 0.8, 2.6, 11.7, 1.6, text, 40, color, bold=True,
          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, font=t.head_font)
    _rect(s, 5.87, 4.3, 1.6, 0.07, t.accent)


def build_pptx(data: dict, titul: list[str], lang: str, path: Path, seed=None) -> Path:
    t = _Theme(random.Random(seed))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)

    _cover(prs, t, data["title"], titul)
    total = len(data["slides"])
    for i, sl in enumerate(data["slides"], 1):
        _content(prs, t, i, total, sl)
    _closing(prs, t, LABELS[lang]["thanks"])

    prs.save(path)
    return path


# ---------------- Word ----------------

def _setup(doc: Document):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    _page_numbers(sec)


def _page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _p(doc, text="", *, bold=False, italic=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
       indent=True, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Cm(1.25) if indent else None
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    return p


def _heading(doc, text, size=14):
    _p(doc, text, bold=True, size=size, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=12, after=6)


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _title_page(doc, kind, data, titul, L):
    univ = titul[0] if titul else ""
    student = titul[1] if len(titul) > 1 else ""
    teacher = titul[2] if len(titul) > 2 else ""
    C = WD_ALIGN_PARAGRAPH.CENTER
    if univ:
        _p(doc, univ.upper(), bold=True, align=C, indent=False)
    for _ in range(7):
        _p(doc, indent=False)
    _p(doc, L[kind], bold=True, size=28, align=C, indent=False)
    _p(doc, f"{L['topic']}: {data['title']}", bold=True, size=16, align=C, indent=False, before=12)
    for _ in range(6):
        _p(doc, indent=False)
    R = WD_ALIGN_PARAGRAPH.RIGHT
    if student:
        _p(doc, f"{L['by']}: {student}", align=R, indent=False)
    if teacher:
        _p(doc, f"{L['checked']}: {teacher}", align=R, indent=False)
    for _ in range(4):
        _p(doc, indent=False)
    _p(doc, str(datetime.now().year), align=C, indent=False)
    _page_break(doc)


def _paragraphs(doc, items):
    for t in items:
        if t.strip():
            _p(doc, t.strip())


def build_docx(kind: str, data: dict, titul: list[str], lang: str, path: Path) -> Path:
    L = LABELS[lang]
    doc = Document()
    _setup(doc)

    if kind == "esse":
        _p(doc, L["esse"], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
        _p(doc, data["title"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
        if titul:
            _p(doc, ", ".join(titul), italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, after=12)
    elif kind == "maqola":
        _p(doc, data["title"].upper(), bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6)
        if titul:
            _p(doc, ", ".join(titul), italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
        if data.get("annotation"):
            p = _p(doc)
            p.add_run(f"{L['annotation']}: ").bold = True
            p.add_run(data["annotation"])
        if data.get("keywords"):
            p = _p(doc, after=12)
            p.add_run(f"{L['keywords']}: ").bold = True
            p.add_run(", ".join(data["keywords"]))
    else:
        _title_page(doc, kind, data, titul, L)
        _heading(doc, L["toc"])
        toc = [L["intro"]]
        for sec in data["sections"]:
            toc.append(sec["heading"])
            toc += [f"    {sub['heading']}" for sub in sec.get("subsections", [])]
        toc += [L["conclusion"], L["refs"]]
        for line in toc:
            _p(doc, line, indent=False)
        _page_break(doc)

    if data.get("introduction"):
        _heading(doc, L["intro"])
        _paragraphs(doc, data["introduction"])
    for sec in data["sections"]:
        _heading(doc, sec["heading"])
        _paragraphs(doc, sec.get("paragraphs", []))
        for sub in sec.get("subsections", []):
            _p(doc, sub["heading"], bold=True, indent=False, before=6, after=3)
            _paragraphs(doc, sub.get("paragraphs", []))
    if data.get("conclusion"):
        _heading(doc, L["conclusion"])
        _paragraphs(doc, data["conclusion"])
    if data.get("references"):
        _heading(doc, L["refs"])
        for i, ref in enumerate(data["references"], 1):
            _p(doc, f"{i}. {ref}", indent=False)

    doc.save(path)
    return path
